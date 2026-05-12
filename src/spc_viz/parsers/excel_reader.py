"""
Top-level Excel orchestration.

Combines :mod:`header_detect`, :mod:`metadata` and :mod:`measurements` to
produce :class:`ParsedFile` instances. Public entry points are
:func:`parse_excel` (single sheet) and :func:`parse_excel_multi` (one or
more data sheets per workbook).
"""

from __future__ import annotations

import io
from typing import IO, Union

import openpyxl
import pandas as pd

from .dimensions import ParsedFile, _safe_str
from .header_detect import _find_data_start, _find_dim_no_cell, _scan_label_rows
from .measurements import (
    _is_non_data_sheet,
    extract_dimension_columns,
    extract_records,
    merge_dimension_groups,
)
from .metadata import build_meta_col_map, coerce_shipment_date, detect_factory
from .openpyxl_patch import _open_strict_ooxml

# Type alias for accepted file inputs (path string or file-like object)
FileOrPath = Union[str, IO[bytes]]


# ---------------------------------------------------------------------------
# Workbook open
# ---------------------------------------------------------------------------


def _open_workbook(file_or_path: FileOrPath) -> tuple[openpyxl.Workbook, str]:
    """Open workbook and return (wb, filename).

    Uses keep_links=False to skip external references.
    If openpyxl returns 0 sheets (strict-OOXML bug), converts the file
    to transitional OOXML in-memory using zipfile XML namespace rewrite.
    """
    if isinstance(file_or_path, str):
        filename = file_or_path.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
        wb = openpyxl.load_workbook(file_or_path, data_only=True, read_only=True, keep_links=False)
        if not wb.sheetnames:
            wb.close()
            wb = _open_strict_ooxml(file_or_path)
    else:
        filename = getattr(file_or_path, "name", "uploaded_file")
        file_or_path.seek(0)
        wb = openpyxl.load_workbook(file_or_path, data_only=True, read_only=True, keep_links=False)
        if not wb.sheetnames:
            wb.close()
            file_or_path.seek(0)
            wb = _open_strict_ooxml(file_or_path)
    return wb, filename


# ---------------------------------------------------------------------------
# Per-sheet orchestrator
# ---------------------------------------------------------------------------


def _parse_single_sheet(
    wb: openpyxl.Workbook,
    sheet_name: str,
    sheet_rows: list,
    dim_no_row: int,
    dim_no_col: int,
    filename: str,
) -> ParsedFile:
    """Parse a single sheet that has already been identified as containing CPK data.

    This is the core parsing logic extracted from parse_excel so it can
    be reused for multi-sheet files.
    """
    label_col = dim_no_col  # column containing row labels
    data_col_start = label_col + 1  # first column of dimension data

    result = ParsedFile(filename=filename, sheet_name=sheet_name)

    # ------------------------------------------------------------------
    # 1. File-level metadata (scan near top for "Part Number", etc.)
    # ------------------------------------------------------------------
    for ri in range(min(5, len(sheet_rows))):
        row = sheet_rows[ri]
        for ci, cell in enumerate(row):
            val = cell.value
            if val is None:
                continue
            s = str(val).strip().lower()
            if "part number" in s and ci + 1 < len(row):
                result.part_number = _safe_str(row[ci + 1].value)
            elif "revision" in s and ci + 1 < len(row):
                result.revision = _safe_str(row[ci + 1].value)
            elif "part description" in s and ci + 1 < len(row):
                result.part_description = _safe_str(row[ci + 1].value)

    # ------------------------------------------------------------------
    # 2. Detect metadata row positions by scanning label column
    # ------------------------------------------------------------------
    label_rows = _scan_label_rows(sheet_rows, label_col, dim_no_row, dim_no_row + 40)

    nominal_row = label_rows.get("nominal")
    tol_max_row = label_rows.get("tol_max")
    tol_min_row = label_rows.get("tol_min")
    usl_row = label_rows.get("usl")
    lsl_row = label_rows.get("lsl")

    # ------------------------------------------------------------------
    # 3. Dimension metadata (scan data columns from data_col_start)
    # ------------------------------------------------------------------
    (
        col_dim_no,
        col_desc,
        col_type,
        col_point,
        col_nominal,
        col_tol_max,
        col_tol_min,
        col_usl,
        col_lsl,
    ) = extract_dimension_columns(sheet_rows, dim_no_row, data_col_start, label_rows)

    # ------------------------------------------------------------------
    # 3b. Merge numbered sub-dimensions (compact format)
    # ------------------------------------------------------------------
    result.dimensions = merge_dimension_groups(
        col_dim_no,
        col_desc,
        col_type,
        col_point,
        col_nominal,
        col_tol_max,
        col_tol_min,
        col_usl,
        col_lsl,
    )

    # ------------------------------------------------------------------
    # 4. Find data start row (auto-detect header + data)
    # ------------------------------------------------------------------
    search_after = max(
        dim_no_row + 10,
        *(v for v in [usl_row, lsl_row, nominal_row, tol_max_row, tol_min_row] if v),
    )
    header_row_idx, data_start_row = _find_data_start(
        sheet_rows, label_col, data_col_start, search_after
    )

    if data_start_row is None:
        # No data rows found; return empty result
        result.data = pd.DataFrame()
        wb.close()
        return result

    # ------------------------------------------------------------------
    # 5. Build metadata column mapping from the header row
    # ------------------------------------------------------------------
    meta_col_map = build_meta_col_map(
        sheet_rows, header_row_idx, label_col, data_col_start, label_rows
    )
    result.meta_columns = list(meta_col_map.keys())

    # ------------------------------------------------------------------
    # 6. Measurement data (from data_start_row onward)
    # ------------------------------------------------------------------
    records = extract_records(sheet_rows, data_start_row, meta_col_map, result.dimensions)
    result.data = pd.DataFrame(records)

    # Convert Shipment Date to datetime if present
    coerce_shipment_date(result.data)

    # ------------------------------------------------------------------
    # 7. Detect factory / site code
    # ------------------------------------------------------------------
    detect_factory(result, filename)

    return result


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def parse_excel(file_or_path: FileOrPath, sheet_name: str = "Raw data") -> ParsedFile:
    """Parse a vendor CPK Excel file and return structured data.

    Auto-detects sheet layout by scanning for "Dim. No." marker cells.
    Works with any sheet name and column/row arrangement.

    Parameters
    ----------
    file_or_path : str or file-like
        Path to an .xlsx file, or an in-memory file object (e.g. from
        Streamlit's file_uploader).
    sheet_name : str
        Hint for which sheet to read. If the exact name or known aliases
        are not found, all sheets are scanned for CPK data layout.

    Returns
    -------
    ParsedFile
    """
    wb, filename = _open_workbook(file_or_path)

    # ----- Auto-detect which sheet(s) contain CPK data -----
    _SHEET_ALIASES: dict[str, list[str]] = {
        "Raw data": ["Raw data", "Raw Data", "raw data", "PP data", "PP"],
        "Data Input": ["Data Input", "data input", "Data input"],
    }
    candidate_sheets: list[str] = []
    # 1. Exact match
    if sheet_name in wb.sheetnames:
        candidate_sheets.append(sheet_name)
    # 2. Aliases
    for alias in _SHEET_ALIASES.get(sheet_name, []):
        if alias in wb.sheetnames and alias not in candidate_sheets:
            candidate_sheets.append(alias)
    # 3. Case-insensitive match
    lower_target = sheet_name.lower()
    for s in wb.sheetnames:
        if s.lower() == lower_target and s not in candidate_sheets:
            candidate_sheets.append(s)
    # 4. ALL remaining sheets (auto-detect mode)
    for s in wb.sheetnames:
        if s not in candidate_sheets:
            candidate_sheets.append(s)

    # Try each candidate; pick first sheet that has "Dim. No." marker
    for sn in candidate_sheets:
        ws = wb[sn]
        rows = list(ws.rows)
        r, c = _find_dim_no_cell(rows)
        if r is not None:
            result = _parse_single_sheet(wb, sn, rows, r, c, filename)
            wb.close()
            return result

    available = ", ".join(wb.sheetnames)
    wb.close()
    raise ValueError(f"No CPK data found in any sheet. Available sheets: {available}")


def parse_excel_multi(
    file_or_path: FileOrPath, sheet_name: str = "Raw data"
) -> list[ParsedFile]:
    """Parse a vendor CPK Excel file and return a list of ParsedFile objects.

    If the requested sheet (e.g. "Raw data") exists, returns a single-element
    list (backward compatible).  If it does not exist, auto-detects ALL data
    sheets by scanning for "Dim. No." markers, skipping known non-data sheets
    (BoxPlotCht*, Histo Pivot, Histo Listbox, Histo Curve).

    This is the preferred entry point for the app layer when a single uploaded
    file may contain multiple data sheets.

    Parameters
    ----------
    file_or_path : str or file-like
        Path to an .xlsx file, or an in-memory file object.
    sheet_name : str
        Preferred sheet name hint (default "Raw data").

    Returns
    -------
    list[ParsedFile]
    """
    wb, filename = _open_workbook(file_or_path)

    # ----- Check for preferred sheet first -----
    _SHEET_ALIASES: dict[str, list[str]] = {
        "Raw data": ["Raw data", "Raw Data", "raw data", "PP data", "PP"],
        "Data Input": ["Data Input", "data input", "Data input"],
    }
    preferred_names: list[str] = []
    if sheet_name in wb.sheetnames:
        preferred_names.append(sheet_name)
    for alias in _SHEET_ALIASES.get(sheet_name, []):
        if alias in wb.sheetnames and alias not in preferred_names:
            preferred_names.append(alias)
    lower_target = sheet_name.lower()
    for s in wb.sheetnames:
        if s.lower() == lower_target and s not in preferred_names:
            preferred_names.append(s)

    # If a preferred sheet exists and has data, return just that (classic path)
    for sn in preferred_names:
        ws = wb[sn]
        rows = list(ws.rows)
        r, c = _find_dim_no_cell(rows)
        if r is not None:
            result = _parse_single_sheet(wb, sn, rows, r, c, filename)
            wb.close()
            return [result]

    # ----- No preferred sheet found: scan all sheets for data -----
    results: list[ParsedFile] = []
    all_sheet_names = list(wb.sheetnames)  # capture before closing
    for sn in all_sheet_names:
        if _is_non_data_sheet(sn):
            continue
        ws = wb[sn]
        rows = list(ws.rows)
        r, c = _find_dim_no_cell(rows)
        if r is not None:
            try:
                parsed = _parse_single_sheet(wb, sn, rows, r, c, filename)
                if parsed.data is not None and len(parsed.data) > 0:
                    results.append(parsed)
            except Exception:
                # Skip sheets that fail to parse
                continue

    wb.close()

    if not results:
        raise ValueError(
            f"No CPK data found in any sheet. Available sheets: {', '.join(all_sheet_names)}"
        )

    return results
